from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
OUT = Path('/Users/shane/Desktop/官方原始数据分析记录.docx')
stats = pd.read_csv(BASE/'temperature_statistics.csv')
quality = json.loads((BASE/'quality_results.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.footer_distance = Cm(.8)
for name, size, bold in [('Normal',11,False),('Title',23,True),('Subtitle',10,False),
                         ('Heading 1',16,True),('Heading 2',13,True),('Heading 3',11,True),
                         ('List Bullet',11,False),('Caption',9,False)]:
    st = doc.styles[name]
    st.font.name = 'Arial Unicode MS'; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
    st.paragraph_format.space_after=Pt(4)
    st.paragraph_format.line_spacing=1.08
    st.paragraph_format.widow_control=True
    for b in st.element.get_or_add_pPr().findall(qn('w:pBdr')): st.element.get_or_add_pPr().remove(b)
    if name.startswith('Heading'):
        st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(12)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Caption'].paragraph_format.space_after=Pt(8)

def p(text, style=None):
    para=doc.add_paragraph(style=style)
    for i,part in enumerate(text.split('**')):
        r=para.add_run(part); r.bold = (i % 2 == 1)
    return para
next_heading_new_page=False
def h(text,level=2):
    global next_heading_new_page
    para=doc.add_heading(text,level=level)
    if next_heading_new_page:
        para.paragraph_format.page_break_before=True
        para.paragraph_format.space_before=Pt(0)
        next_heading_new_page=False
    return para
def bullet(text): return p(text,'List Bullet')
def page():
    global next_heading_new_page
    next_heading_new_page=True
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    border=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
    t._tbl.tblPr.append(border)
    for j,v in enumerate(headers): t.rows[0].cells[j].text=v
    for values in rows:
        cells=t.add_row().cells
        for j,v in enumerate(values):cells[j].text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','85'),('bottom','85'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            fill='173B50' if i==0 else ('F1F5F7' if i%2==0 else 'FFFFFF')
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);tcpr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(9.5);r.font.bold=(i==0)
                    r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def figure(name,caption):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(3)
    para.paragraph_format.keep_with_next=True
    pic=para.add_run().add_picture(str(BASE/name),width=Cm(15.2))
    pic._inline.docPr.set('descr',caption)
    p(caption,'Caption')
def value(setpoint,metric,key='mean'):
    return float(stats[(stats.setpoint_C==str(setpoint)) & (stats.metric==metric)].iloc[0][key])
def filelink(label,path):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
    rel=doc.part.relate_to(Path(path).resolve().as_uri(),RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');pr.append(OxmlElement('w:u'))
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);link.append(r);para._p.append(link)
    return para

doc.core_properties.title='官方原始数据分析记录'
doc.core_properties.subject='TechArena 2026 Challenge 2 官方原始数据的统计与质量检查'
doc.core_properties.author='项目组'
doc.core_properties.keywords='LFP SOH 温度 原始数据 数据质量'
p('官方原始数据分析记录','Title')
p('TechArena 2026 Challenge 2　｜　版本 1.0　｜　更新于 2026 年 10 月 6 日','Subtitle')
p('本记录用于持续汇总官方原始数据的字段含义、分布规律、缺失与异常，以及对后续分析的影响。当前已完成温度章节；后续电流、电压、累计 Ah 等总结可作为独立章节追加，并在更新记录中登记范围和结论变化。')
table(['更新日期','版本','本次内容'],[['2026-10-06','1.0','新增温度统计、缺失检查、异常筛查及三张图表']],[3.1,1.7,12])
h('第一章 温度数据',1)
p('本章依据官方 data/operation/ 全部 13 个运行文件，合计 **1,572,894 行**，记录时间为 **2025-01-19 19:10:56 至 2025-09-11 18:07:11**。分析覆盖全部已公开运行数据，包含 CK7 后记录。')
h('主要结论')
bullet('温箱设定 25°C 时，实测平均为 **25.63°C**；设定 45°C 时，实测平均为 **45.94°C**，分别平均高出 **0.63°C、0.94°C**。')
bullet('全体记录的最高减最低温差平均为 **1.14°C**；平均比最低高约 **0.66°C**，最高比平均高约 **0.48°C**。')
bullet('已有行的温度字段缺失和大小关系错误均为 **0**，但存在未提供运行段和采样中断。')
bullet('按本章探索规则，共 **67 行**需要复核，占 **0.00426%**；这不等于已经确认 67 行坏数据。')
h('温度字段与统计口径')
table(['字段','含义'],[
['chamber_temperature_C','温箱设定值，只有 25°C、45°C 两档'],
['temp_mean_C','电芯温度传感器读数的平均值'],
['temp_min_C','同一时刻传感器读数的最低值'],
['temp_max_C','同一时刻传感器读数的最高值']],[5.5,11.3])
p('“标称温度”在本章指温箱设定值，不是温箱空气温度的独立实测值。各温度传感器的数量、位置、精度及独立通道未提供，不能将最大或最小温度归属于某一枚特定电芯。',None)

page()
h('平均温度与温箱设定值')
p('25°C 档有 **529,012 行**，45°C 档有 **1,043,882 行**。先逐行计算温差，再按记录行等权汇总；“90% 区间”为第 5 至第 95 百分位，不是预测置信区间。')
table(['指标','设定 25°C','设定 45°C'],[
['实测平均温度的总体均值','25.627°C','45.939°C'],
['实测平均减设定值的平均偏差','+0.627°C','+0.939°C'],
['偏差中位数','+0.56°C','+0.91°C'],
['偏差的 90% 区间','+0.19 至 +1.31°C','+0.50 至 +1.51°C'],
['实测平均温度的全部范围','24.89 至 27.49°C','44.10 至 47.61°C'],
['实测平均高于设定值的行占比','99.86%','99.995%']],[7.6,4.6,4.6])
p('实测平均温度通常高于设定值，但偏差并不恒定。充放电发热、温箱调节、传感器位置或校准差异都是可能解释；当前数据不能将它们分开，也不能将偏差全部认定为测量误差。')
p('按电流描述性分组，25°C 档充电／放电平均偏差为 +0.729／+0.500°C；45°C 档为 +1.032／+0.827°C。充电定义为 >0.5 A，放电为 <-0.5 A。该比较没有匹配 SOC、时间和静置状态，不能视作独立发热效应。')
figure('temperature_timeline.png','图 1　官方运行温度及传感器温差。来源为 2025 年 1至9月运行数据，每 30 分钟按已有样本取平均；空白不插值。浅色区域是最低与最高温度的时间窗均值，不是原始极值包络。')
p('逐段偏差也在变化：45°C 档第 1 段为 +1.071°C，第 13 段为 +0.789°C。可见单一固定修正量不能描述全部片段；变化原因和老化贡献仍需另外验证。')

page()
h('平均温度与最高最低温度')
table(['指标','设定 25°C','设定 45°C'],[
['最低温度的总体均值','24.949°C','45.287°C'],
['平均温度的总体均值','25.627°C','45.939°C'],
['最高温度的总体均值','26.104°C','46.415°C'],
['平均减最低的平均差','0.678°C','0.652°C'],
['最高减平均的平均差','0.477°C','0.476°C'],
['最高减最低的平均差','1.155°C','1.127°C'],
['最高减最低的 90% 区间','0.90 至 1.40°C','0.90 至 1.40°C'],
['最高减最低的最大值','2.20°C','3.60°C']],[7.6,4.6,4.6])
p('两档的通常温差接近，45°C 档没有表现出更大的平均传感器温差。全体最高减最低平均为 **1.137°C**，两档的第 1 至第 99 百分位区间均为 **0.8 至 1.5°C**。由于大量读数重复，区间包含比例不必恰好等于 98%。')
p('平均温度并非最高与最低的简单中点：全体平均值比两端中点高约 **0.092°C**，略偏向较高温一侧。这是传感器统计关系，不能定位具体偏热电芯。')
figure('temperature_distributions.png','图 2　平均温度偏差与最高减最低温差的分布。来源为全部官方运行记录，25°C、45°C 分别计算档位内部样本百分比，未按日历时间补齐缺失记录。')
h('记录中的实测取值范围')
p('平均温度为 **24.89 至 47.61°C**，最低通道为 **23.9 至 47.0°C**，最高通道为 **25.2 至 48.3°C**。没有明显零值占位或巨大数值，但取值范围合理并不证明传感器校准准确。')

page()
h('字段缺失与异常筛查')
p('四个温度字段共 **6,291,576 个数值位置**。以下检查针对已提供记录；时间覆盖缺口另列下一页。')
table(['检查项','结果'],[
['空值或 NaN','0 个'],['任一温度字段缺失的记录','0 行'],
['正负无穷等非有限值','0 个'],['三个实测温度字段中的零值','0 个'],
['设定值不是 25°C 或 45°C','0 行'],['不满足最低 ≤ 平均 ≤ 最高','0 行']],[12,4.8])
p('本章采用两条**探索筛查规则**，不是官方报警标准或仪器精度阈值：')
bullet('最高减最低 **>2°C**：**63 行**，占 0.00401%；其中 >3°C 有 18 行。')
bullet('同文件相邻有效记录的任一实测温度变化 **>0.5°C**：**9 行**。只比较时间间隔 >0 且 ≤30 秒的记录；命中的间隔均为 10 秒，全部来自最低温度通道，变化最大 0.8°C。')
p('两类重叠 5 行，合并后为 **67 行，占 0.00426%**。若将突跳阈值改为 >1°C 或 >2°C，均为 0 行。67 行分布于 1 月 22 日（10 行）、4 月 18 日（8 行）、7 月 11 日（31 行）、7 月 14 日（18 行）。')
figure('largest_spread_episode.png','图 3　第 11 段最大温差附近的原始采样。2025-07-11 13:52:02：设定 45°C，最低 42.40°C，平均 44.49°C，最高 46.00°C，温差为全体最大值 3.60°C。')
p('附近数分钟内最低温度多次下降和恢复，最高温度变化较小，平均值也下降。局部冷却、传感器接触或采集问题等解释需设备日志和独立通道核对。它不是单个孤立大数值，当前应保留并标记，不能直接判为传感器损坏或删除。')

page()
h('时间覆盖与重复时间戳')
p('**整段缺失：**官方第 3 段未提供，说明标注覆盖 2025 年 2 月 19 日至 3 月 11 日。另有月度 checkup 等段间暂停，不应全部称为传感器丢样。')
p('**文件内部中断：**以相邻间隔 >30 秒筛查，共 6 处。间隔合计约 20.81 小时；每处减去一份常规 10 秒间隔后，超出约 **20.79 小时**。最长 59,909 秒，约 16 小时 38 分钟。该数字不含第 3 段及文件间隔，不能换算为精确的全实验缺失行数。')
g=pd.read_csv(BASE/'within_file_recording_gaps.csv')
table(['运行段','前一条记录时间','后一条记录时间','间隔'],[
[str(int(r.segment)),str(r.previous_timestamp)[5:],str(r.timestamp)[5:],f'{int(r.gap_seconds):,} 秒'] for r in g.itertuples()
],[1.5,5.7,5.7,3.9])
p('**重复时间戳：**181 组，相当于 181 条超出唯一时间戳数量的记录。其中 119 组至少一个实测温度不同，同一时间戳最大温度差为 0.20°C。这属于采样与时钟对应关系待核对，尚不是已证明的温度坏值。')
p('**参考测量温度：**CK0 参考放电文件没有温度列，因此无法核对初始容量测量时的实际温度；这是未提供字段，与运行表空值不同。')
h('统计权重与使用边界')
p('主统计按记录行等权，没有去重或跨段补值。时间加权检查采用同文件下一条记录时间差，仅 0<间隔≤30 秒时计权，零间隔和长缺口不计。两档平均偏差与按行结果之差均 <0.00013°C，主要结论一致。')
p('本章覆盖 CK7 之后的公开数据。用于早期 CK 预测时，必须按截止时刻重算事件和统计，不能直接将本章全量均值作为早期参考。')
h('对后续分析的建议')
bullet('分别保留实测平均温度、温箱设定温度及最高减最低温差。偏差不是纯测量误差，不作统一常数校正。')
bullet('标记 67 行及其附近片段，后续比较保留与屏蔽的敏感性；长缺口和缺失段不直接线性插值。')
bullet('本章只说明温度与数据质量，不能证明温度特征改善容量预测或候选片段属于电池故障；CK1至CK7 真值仍隐藏。')

page()
h('证据与复现入口')
p('以下路径均相对于项目文件夹“Arena 阶段 2 官方材料”。点击链接可打开本机文件；文档中的图表已嵌入，移动 Word 文件不影响查看图表。')
for label,relative in [
('官方数据字段说明','data/DATA_DESCRIPTION.md'),
('温度分析完整文字记录','官方温度数据分析总结_20261006.md'),
('逐温度档位统计与百分位','research/official_temperature_review_20261006/temperature_statistics.csv'),
('逐运行段温度统计','research/official_temperature_review_20261006/segment_statistics.csv'),
('67 行候选及相邻温度记录','research/official_temperature_review_20261006/temperature_review_candidates.csv'),
('采样中断清单','research/official_temperature_review_20261006/within_file_recording_gaps.csv'),
('质量检查结果','research/official_temperature_review_20261006/quality_results.json'),
('完整结果及 13 个输入文件 SHA256','research/official_temperature_review_20261006/analysis_results.json'),
('温度分析脚本','research/official_temperature_review_20261006/analyze_temperature.py')]:
    filelink(label,ROOT/relative)
p('上述统计来自 2026 年 10 月 6 日对全部公开运行文件的分析。按运行段核对总行数为 1,572,894；候选清单核对 63+9−5=67 行，来源文件和时间戳组合无重复。')
h('后续章节的记录方式',1)
p('后续原始数据总结沿用 Word 的“标题 1”添加第二章、第三章等，章节内部使用“标题 2”组织内容，以便通过导航窗格查阅。当前不预填未分析字段的结论。')
bullet('**范围和字段：**记录来源文件、时间范围、行数、字段含义、单位及已知限制。')
bullet('**统计结果：**列出总体和必要分组的均值、分位数与极值，并说明权重、筛选与图表口径。')
bullet('**缺失与异常：**区分字段空值、整段缺失、时间中断和筛查候选；明确阈值、数量、占比和重叠。')
bullet('**结论与证据：**区分直接观测与解释假设，保存原始行、统计清单和复现入口。')
bullet('**更新登记：**新增章节或修正旧结论时，更新首页日期、版本及更新记录，注明被修正的内容和依据。')

f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('官方原始数据分析记录　｜　第 ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
f.add_run(' 页').font.size=Pt(9)
update_fields=OxmlElement('w:updateFields');update_fields.set(qn('w:val'),'true');doc.settings.element.append(update_fields)
doc.save(OUT)
print(OUT)
