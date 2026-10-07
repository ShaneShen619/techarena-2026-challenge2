"""Build the Chinese decision report DOCX from its audited Markdown source."""
import re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TASK=Path(__file__).resolve().parents[1]
src=TASK/'outputs/LFP_SOH_官方基线修复与下一步决策报告.md'
dest=src.with_suffix('.docx')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.58)
sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)

def set_font(style,size,bold=False):
    style.font.name='PingFang SC';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr();fonts=rpr.rFonts
    if fonts is None:
        fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'),'PingFang SC')

set_font(doc.styles['Normal'],10)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1.13
set_font(doc.styles['Title'],18,True)
doc.styles['Title'].paragraph_format.space_after=Pt(11)
tp=doc.styles['Title'].element.get_or_add_pPr()
for border in tp.findall(qn('w:pBdr')):tp.remove(border)
set_font(doc.styles['Heading 1'],12.5,True)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True

def add_inline(p,line):
    chunks=re.split(r'(\*\*.*?\*\*|`[^`]+`)',line)
    for c in chunks:
        if not c:continue
        if c.startswith('**') and c.endswith('**'):p.add_run(c[2:-2]).bold=True
        elif c.startswith('`') and c.endswith('`'):
            r=p.add_run(c[1:-1]);r.font.name='Menlo';r.font.size=Pt(9)
        else:p.add_run(c)

lines=src.read_text().splitlines();title_done=False
for line in lines:
    s=line.strip()
    if not s:continue
    if s.startswith('# '):
        p=doc.add_paragraph(style='Title');p.add_run(s[2:]);title_done=True
        pp=p._p.get_or_add_pPr();bd=pp.find(qn('w:pBdr'))
        if bd is not None:pp.remove(bd)
    elif s.startswith('## '):
        doc.add_paragraph(s[3:],style='Heading 1')
        if s[3:]=='B4 至 B6 官方接口和决策':
            matrix=pd.read_csv(TASK/'outputs/baseline_comparison_matrix.csv')
            t=doc.add_table(rows=1,cols=5)
            t.autofit=False
            widths=[.55,1.35,1.35,1.35,1.35]
            labels=['CK','原样示例 %','因果等价 %','起点门 %','CK0 常数 %']
            for i,(cell,label) in enumerate(zip(t.rows[0].cells,labels)):
                cell.text=label;cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'DCE7F3');tcPr.append(shd)
                for r in cell.paragraphs[0].runs:r.bold=True;r.font.size=Pt(8.5)
            for row in matrix.itertuples():
                vals=[row.checkup,f'{row.original_pct:.3f}',f'{row.SOH_est_pct:.3f}',f'{row.start_pct:.3f}',f'{row.CK0_constant_pct:.3f}']
                cells=t.add_row().cells
                for i,(cell,val) in enumerate(zip(cells,vals)):
                    cell.text=str(val);cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for p in cell.paragraphs:
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:r.font.size=Pt(8.5)
            tblPr=t._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement(f'w:{edge}');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tblPr.append(borders)
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            p=doc.add_paragraph('表 1  官方无标签预测对照。所有容量误差均不可计算；起点门为 B1 后探索假设。')
            p.style=doc.styles['Normal'];p.runs[0].italic=True
    else:
        p=doc.add_paragraph();add_inline(p,s)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('LFP SOH 阶段二研究  |  2026 年 9 月 30 日  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='LFP SOH 官方基线修复与下一步决策报告'
doc.core_properties.subject='阶段二官方基线因果审计与质量门研究'
doc.save(dest)
print(dest)
