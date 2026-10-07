"""Deterministically typeset the validated Markdown report with python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

TASK=Path(__file__).resolve().parents[1]
SOURCE=TASK/'outputs/LFP_SOH_双方案设计验证与比较报告.md'
OUTPUT=SOURCE.with_suffix('.docx')

def set_font(style,name,size,bold=False):
    style.font.name=name;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(20,27,32)
    rpr=style.element.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
    for attr in ('ascii','hAnsi','eastAsia','cs'):rfonts.set(qn('w:'+attr),name)

def add_hyperlink(p,label,url):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'275E7A');properties.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
    run.append(properties);node=OxmlElement('w:t');node.text=label;run.append(node)
    link.append(run);p._p.append(link)

def add_inline(p,text):
    # Bold, monospace, and Markdown links are represented as styled runs.
    token=re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))')
    parts=token.split(text)
    for part in parts:
        if not part:continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]);r.bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Menlo';r.font.size=Pt(8.2)
        elif part.startswith('[') and '](' in part:
            label,url=part[1:-1].split('](',1)
            if url.startswith('http'):
                add_hyperlink(p,label,url)
            else:p.add_run(label)
        else:p.add_run(part)

def add_table(doc,lines):
    data=[]
    for line in lines:
        cells=[c.strip() for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?',c or '') for c in cells):continue
        data.append(cells)
    if not data:return
    n=max(len(row) for row in data)
    table=doc.add_table(rows=0,cols=n)
    table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=True
    for idx,values in enumerate(data):
        cells=table.add_row().cells
        for j,val in enumerate(values):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(0)
            add_inline(p,val)
            for run in p.runs:run.font.size=Pt(8.0 if n>=4 else 8.5)
            if idx==0:
                for run in p.runs:run.bold=True
                tcPr=cells[j]._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EAF0F2');tcPr.append(shd)
        if idx==0:
            trPr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trPr.append(repeat)
    # These report tables are small enough to keep intact. In particular,
    # avoid leaving a header plus one cell row below a figure at page end.
    for row in table.rows[:-1]:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next=True
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def add_figure(doc,line):
    match=re.fullmatch(r'!\[(.+)\]\((.+)\)',line)
    if not match:return
    caption,rel=match.groups();path=SOURCE.parent/rel
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
    run=p.add_run();shape=run.add_picture(str(path),width=Inches(6.05))
    shape._inline.docPr.set('descr',caption)
    c=doc.add_paragraph('图  '+caption)
    c.style='Caption';c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after=Pt(9)

def main():
    lines=SOURCE.read_text().splitlines()
    doc=Document()
    sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.62)
    sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    set_font(doc.styles['Normal'],'PingFang SC',9.4)
    set_font(doc.styles['Title'],'PingFang SC',20,True)
    set_font(doc.styles['Heading 1'],'PingFang SC',13.5,True)
    set_font(doc.styles['Heading 2'],'PingFang SC',11.0,True)
    set_font(doc.styles['Caption'],'PingFang SC',8.4)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    # The built-in Title style may carry a colored bottom border in the
    # renderer's default template; the requested report title uses whitespace.
    title_ppr=doc.styles['Title'].element.get_or_add_pPr()
    for border in title_ppr.findall(qn('w:pBdr')):title_ppr.remove(border)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_inline(p,line[2:])
            p.paragraph_format.space_after=Pt(11)
            ppr=p._p.get_or_add_pPr()
            for border in ppr.findall(qn('w:pBdr')):ppr.remove(border)
        elif line.startswith('## '):
            doc.add_heading(line[3:],1)
        elif line.startswith('### '):
            doc.add_heading(line[4:],2)
        elif line.startswith('!['):
            add_figure(doc,line)
        elif line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip());i+=1
            add_table(doc,block);continue
        else:
            p=doc.add_paragraph();add_inline(p,line)
            if line.startswith('2026 年'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        i+=1
    doc.core_properties.title='LFP 电池 SOH 第二阶段双方案设计验证与比较报告'
    doc.core_properties.subject='路线 A 与路线 B 的设计、实测代理验证和边界'
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__=='__main__':main()
