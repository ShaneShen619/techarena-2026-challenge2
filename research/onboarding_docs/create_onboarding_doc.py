"""Create a teammate-facing guide to the Phase 2 written material."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "LFP电池SOH第二阶段_新队友资料导读_20260928.docx"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.75)
sec.left_margin, sec.right_margin = Cm(2.05), Cm(2.05)
sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)

INK = RGBColor(30, 43, 54)
BLUE = RGBColor(29, 72, 105)
MUTED = RGBColor(96, 110, 121)
PALE = "EAF1F6"


def set_font(style, size, bold=False, color=INK, space_after=5):
    s = doc.styles[style]
    s.font.name = "Songti SC"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = color
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    s.paragraph_format.space_after = Pt(space_after)
    return s


set_font("Normal", 9.7, space_after=5)
set_font("Title", 21, True, INK, 8)
set_font("Subtitle", 10, False, MUTED, 10)
set_font("Heading 1", 13.2, True, BLUE, 7)
set_font("Heading 2", 10.7, True, BLUE, 4)
set_font("List Bullet", 9.7, space_after=2)
set_font("List Number", 9.7, space_after=2)
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)
for name in ("Heading 1", "Heading 2"):
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles["Normal"].paragraph_format.line_spacing = 1.15


def p(text="", *, style=None, bold_lead=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def h(text, level=1):
    doc.add_heading(text, level=level)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    if widths:
        for i, width in enumerate(widths):
            t.columns[i].width = Cm(width)
    for i, label in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = label
        shade(cell, PALE)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = BLUE
            r.font.size = Pt(8.8)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.name = "Songti SC"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
                    run.font.size = Pt(8.5)
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


title = doc.add_paragraph(style="Title")
title.add_run("LFP 电池 SOH 第二阶段资料导读")
sub = doc.add_paragraph(style="Subtitle")
sub.add_run("给刚加入项目的队友  |  截至 2026 年 9 月 28 日  |  建议先读本页再查原文")

h("先用一分钟理解项目")
p("我们要根据一组四串、每芯标称 102 Ah 的 LFP 电池运行记录，在 CK0—CK7 八个日期估计电池组的容量 SOH。定义是 100 × 按约 5.1 A 放电至组压 11.2 V 的容量 / 102 Ah。只有 CK0=100.41 Ah（98.44%）公开；CK1—CK7 的真实容量隐藏，每个点只能用它之前的运行信息。第一阶段是长期寿命预测，第二阶段是运行中的当前 SOH 估计。[1][2]")
p("这两天的文档经历了“读规则和文献 → 做首选/备选原型 → 修温度与容量映射 → 重新评估 → 扩展多路线并严格审计”的过程。最新结论是：研究和接口验证已经完成一轮闭环，预设的高精度目标没有达到，也没有官方隐藏点的真实误差。不能把任何本地代理数据上的成绩称为官方四串成绩。[5][8][9][10]")

table(["最关键的三件事", "对新队友的含义"], [
    ("只有 CK0 容量真值", "本地能验证格式、时间因果和信号规律；不能计算 CK1—CK7 官方容量 MAE。"),
    ("数据和标签有不同层级", "P1 六芯是单芯高倍率放电代理；Che 有异协议充电容量；TU 适合阻抗；合成只检验机制。"),
    ("研究包不是最终提交", "已有两个通过接口硬测试的研究候选包，但项目根 my_model/__init__.py 仍启用官方 ExampleModel；尚需团队决定和完成最终提交物。"),
], [5.0, 11.8])

h("推荐阅读顺序")
table(["用时", "先读什么", "读完应该知道"], [
    ("10 分钟", "本导读 → research/phase2_method_exploration/outputs/EXECUTIVE_SUMMARY.md", "题目、最新状态、性能目标为何未达。"),
    ("再用 20 分钟", "data/DATA_DESCRIPTION.md → submission_instructions_2026.md → 最新多路线报告", "输入和输出、哪些是真值、已验证和未验证的结论。"),
    ("若要接手技术", "文献综述 → 双方案报告 → 温度报告与复评 → M9_FINDINGS.md → REPRODUCE.md", "方案为何改变、历史成绩为何不能直接复用、如何核验。"),
], [2.2, 8.2, 6.4])

h("一 资料地图  每份材料在解决什么问题")
p("路径如无特别说明，均相对桌面文件夹“Arena 阶段 2 官方材料”。桌面独立 Word 文件会特别注明。旧报告保留过程证据，但判断当前状态时以最新验收和审计为准。")
table(["材料", "作用与阅读提醒"], [
    ("桌面《LFP电池SOH第二阶段会议总结.docx》", "启动会口径：两阶段关系、任务、评分、提交、现场问答；包含另一赛题和出行信息，只取 Topic 1 电池部分作为本项目依据。"),
    ("桌面《LFP电池SOH第二阶段项目任务与实施说明.docx》", "把官方题目转成项目执行清单、字段表和提交清单。其“模型尚未实现”是启动时快照，已被后续研究进度更新。"),
    ("官方 CHALLENGE2_DESCRIPTION.pdf、README.md、submission_instructions_2026.md、data/DATA_DESCRIPTION.md", "任务定义、数据字段、接口和提交边界的原始依据。上传入口、截止时间等带 TODO 的条目仍须核实。"),
    ("research/phase2_literature/outputs/LFP_SOH_第二阶段文献综述与技术路线.docx", "比较部分充电、ICA/DVA、脉冲、温度、状态模型等，提出最初的首选 A 与备选 B；这是研究假设，不是后续实验的最终结论。"),
    ("research/phase2_validation/outputs/LFP_SOH_双方案设计验证与比较报告.docx", "真正实现首选 A（单芯窗口比）和备选 B（受可辨识性门控的状态模型），给出 18 个单芯开发目标及官方无标签回放。"),
    ("research/phase2_temperature_improvement/outputs/方案A_温度修复与提升报告.docx", "在 P1 六芯 180 点深充开发面板上修参考、事件和温度，宏 MAE 1.039 pp；该输入条件不同于官方浅充。"),
    ("research/phase2_temperature_improvement/reassessment_20260928/REEVALUATION.md", "回头检查温度贡献、整段深充 Ah 的主导作用和官方备用分支风险；是转向严格浅充验证的原因。"),
    ("research/phase2_method_exploration/outputs/EXECUTIVE_SUMMARY.md、LFP_SOH_多路线设计验证与比较报告.md", "最新主结论。M0—M10 逐一路线、消融、失败和适用范围；判断项目现在到了哪里应优先引用这里。"),
    ("research/phase2_method_exploration/outputs/MEASUREMENT_PLAN.md", "若要真正确认官方容量精度，需要如何采独立四串容量、真实浅充、脉冲、温度和均衡。它是建议方案，尚未执行。"),
    ("research/phase2_next_round/START_PROMPT.md", "下一轮 R1—R6 的待启动研究任务书；它不是已完成结果。"),
], [8.1, 8.7])

h("二 数据与证据不能混用")
table(["来源", "有什么", "能支持什么  不能支持什么"], [
    ("官方 Challenge 2", "一组四串；13 段运行，通常 10 秒采样；CK0 容量与放电曲线。", "可检查前缀、事件、未来电压和模型运行。CK1—CK7 隐藏，无法本地量化容量误差。"),
    ("P1 第一阶段", "六个物理单芯，跨寿命有放电容量。", "能做留整芯开发验证；标签为 0.5C/1C 至单芯 2.5 V，不等于四串 C/20 组压截止容量。"),
    ("Che Dataset 3", "11 个单芯的部分充电曲线与充电容量。", "可看部分充电形状；本地包缺完整原始时间、电流、电压、温度时序，协议也不等于官方标签。"),
    ("TU Darmstadt", "28 个八串现场系统，电压、电流、温度和阻抗相关信号。", "可研究工况归一化阻抗与异常；没有逐次真实容量标签，BMS SOC 不是 SOH。"),
    ("合成场景", "人为设定的逐芯容量、SOC、阻抗和截止。", "可反证公式与机制；不能替代真实组容量测量。"),
], [2.9, 5.4, 8.5])
p("阅读任何结果时先问三个问题：目标容量来自哪里？物理独立单位是单芯还是电池组？模型在预测时是否真的能看到这些输入？六芯的 180 个目标点不是 180 个独立电芯。", bold_lead="阅读任何结果时先问三个问题：")

h("三 方案为什么一路变化")
table(["阶段", "当时学到的事", "当前如何使用"], [
    ("文献与最初路线", "部分充电可能含健康信息；LFP 平台让容量、SOC、极化混淆。首选 A 是对齐单芯窗口；备选 B 是受约束状态估计。", "当作假设地图和原论文入口，不能照搬论文误差。"),
    ("双方案实跑", "A 在六芯 18 目标的单芯代理上为 10.57 pp，B 为 15.27 pp；B 在官方运行零次可靠容量更新。两者接口可运行。", "A 是早期有条件候选，B 是可辨识性/保守回退。官方八点输出均无容量真值。"),
    ("温度与深充修复", "P1 六芯 180 点深充开发面板的修复模型 1.039 pp；温度边际改善约 0.098 pp，整段 CC 充入 Ah 是更强信号。", "只说明特定深充代理任务有用。官方日常多为约 21 Ah 浅充，实际均进入备用分支。"),
    ("复评与严格浅充", "发现旧浅充 v1.4 先用完整充电事件挑样，再裁到 3.50 V，选样用到终点后的信息；旧较好成绩退役。", "以重新扫原始前缀并重训的 v1.5 为当前公平主比较。"),
    ("最新多路线研究", "在 v1.5 六芯 D1（深充记录裁剪出的浅充代理）上，最强简单基线 3.755 pp，TCN 微调集成 2.851 pp，但最差芯 5.572 pp、最大单点 17.175 pp。", "相对改善 24.08%，但预设宏平均≤1 pp、最差芯≤2 pp、最大单点≤5 pp 的目标均未达；波形增量未获证实。"),
], [3.1, 8.1, 5.6])
p("这些数值来自不同目标数、可见输入和标签协议，不能按大小直接排名。尤其 1.039 pp 属深充开发，2.851 pp 属严格浅充裁剪开发，10.57 pp 属早期 18 点原型；它们都不是官方 CK1—CK7 的真实误差。")

h("四 最新研究的结论  哪些路线留下来")
table(["路线", "目前能说的结论"], [
    ("浅充窗口与非线性模型", "可在 P1 代理上降低均值误差，但终点、温度和最差芯脆弱。去充电电压后 TCN 几乎不退化，波形独立容量信息未证实。"),
    ("温度", "旧第一阶段温度面迁到官方设备约有 10°C 偏差；合法片段温度＋标量 MLP 在 D1 为 2.933 pp，但不能分离温度因果作用与芯/工况身份。"),
    ("重复标准脉冲", "官方 41 段脉冲已认证，后续逐芯电压能较好预测；约 10 Ah 固定脉冲不是容量标签。"),
    ("状态估计与四串截止", "官方浅平台片段对逐芯容量不可辨，148 次拟合都撞边界，门控不更新。四串截止方程可重建 CK0、通过合成机制测试，但真实逐芯状态尚未证实。"),
    ("TU、自监督与融合", "TU 可做阻抗/异常；TCN 微调的代理数值改善不能归因于健康波形；融合未稳定超过最佳单路。"),
    ("当前研究候选", "CK0 常数包与旧浅充备用分支包均通过干净提取接口验证和九项硬测试。它们是可审计研究对照，不是已证明准确的官方提交方案。"),
], [4.0, 12.8])

h("五 接手时先做什么")
for line in [
    "先确认分工：研究论文与假设、数据/事件审计、模型代码、独立验证、报告与提交各由谁负责；不要让多个人同时改同一候选包或验收文件。",
    "先读 research/phase2_method_exploration/README.md 与同目录 REPRODUCE.md，运行只读/派生核验，对齐 correctness_pass=true、research_closed=true、d1_target_met=false 的含义；检查固定面板哈希和候选包验证收据。",
    "打开最新报告与 outputs/official_diagnostics.csv，逐点查看 CK1—CK7 的输入覆盖、备用分支、跳变和回退；只讨论模型行为，勿推断隐藏点真值。",
    "若继续建模，按下一轮 START_PROMPT 的 R1—R6 顺序，先冻结同预算、留物理实体和严格前缀协议，再做标量先验、浅充波形、脉冲、状态、外部数据和融合的反证。",
    "若能推动新测量，优先执行 MEASUREMENT_PLAN：新独立四串组、真实 15/20/30 Ah 浅充、25/45°C 配对、标准脉冲与静置、重复 5.1 A 至组压首次 11.2 V 容量；封存组只在候选锁定后一次确认。",
    "提交前单独核实官方截止时间和上传入口；当前根目录 ActiveModel 仍指向 ExampleModel。团队选定正式方案后才切换，并按官方要求完成依赖、README、报告、两阶段 PPT、干净解压验证和最终 ZIP。",
]:
    p(line, style="List Number")

h("六 查数值和代码时去哪儿")
table(["想查的问题", "入口"], [
    ("最新项目结论", "research/phase2_method_exploration/outputs/EXECUTIVE_SUMMARY.md 与完整报告"),
    ("每条路线为何有效或失败", "research/phase2_method_exploration/notes/M1_FINDINGS.md 至 M9_FINDINGS.md"),
    ("最终验收字段", "research/phase2_method_exploration/outputs/acceptance_results.json"),
    ("逐点预测与分层比较", "research/phase2_method_exploration/outputs/predictions_long.csv、method_comparison.csv"),
    ("官方各 CK 实际用了什么", "research/phase2_method_exploration/outputs/official_predictions.csv、official_diagnostics.csv"),
    ("错误、版本和审查", "research/phase2_method_exploration/notes/DECISIONS.md、FAILURES.md、FINAL_AUDIT.md、READER_TEST.md"),
    ("怎么重新核验", "research/phase2_method_exploration/REPRODUCE.md、notes/EXPERIMENT_INDEX.csv"),
], [5.0, 11.8])

h("术语速查")
p("SOH：健康状态，这里特指按指定低倍率放电协议测得的容量与 102 Ah 标称容量之比。pp：SOH 百分点，误差从 90% 到 88% 是 2 pp。CK：月度检查点。严格前缀：预测某时刻时，事件筛选、特征和模型状态都不使用之后的信息。D0/D1：P1 深充开发/从深充裁剪出的浅充开发代理；D2：尚缺的独立兼容四串容量确认；D3：官方无标签回放；S：合成机制。")

doc.add_page_break()
h("交接会上要定的五件事")
table(["需要定的事", "为什么现在要定"], [
    ("最终模型负责人及版本", "项目根仍运行 ExampleModel；研究候选尚未被选定为正式提交。"),
    ("新数据是否能取得", "独立四串 C/20 容量真值是当前最重要的证据缺口，决定能否确认准确度。"),
    ("下一轮优先路线与停机标准", "R1—R6 的任务书已写，需按冻结输入和实体留出逐路做，避免反复调六芯开发集。"),
    ("复现与代码审查分工", "一人维护候选，一人按 REPRODUCE/逐点输出核验因果、标签、单位和失败案例。"),
    ("提交与展示事项", "核实实际截止/上传入口，安排正式 README、报告、两阶段 PPT、依赖和干净解压验证。"),
], [4.8, 12.0])

h("依据文件")
sources = [
    "[1] 桌面《LFP电池SOH第二阶段会议总结.docx》；官方 CHALLENGE2_DESCRIPTION.pdf。",
    "[2] data/DATA_DESCRIPTION.md；submission_instructions_2026.md；项目根 README.md。",
    "[3] 桌面《LFP电池SOH第二阶段项目任务与实施说明.docx》（早期状态快照）。",
    "[4] research/phase2_literature/outputs/LFP_SOH_第二阶段文献综述与技术路线.docx。",
    "[5] research/phase2_validation/outputs/LFP_SOH_双方案设计验证与比较报告.docx。",
    "[6] research/phase2_temperature_improvement/outputs/方案A_温度修复与提升报告.docx。",
    "[7] research/phase2_temperature_improvement/reassessment_20260928/REEVALUATION.md。",
    "[8] research/phase2_method_exploration/outputs/LFP_SOH_多路线设计验证与比较报告.md。",
    "[9] research/phase2_method_exploration/outputs/acceptance_results.json；notes/FINAL_AUDIT.md。",
    "[10] research/phase2_method_exploration/outputs/MEASUREMENT_PLAN.md；research/phase2_next_round/START_PROMPT.md。",
]
for line in sources:
    para = p(line)
    para.paragraph_format.space_after = Pt(2)
    for run in para.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("项目资料导读  ·  2026-09-28  ·  第 ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
