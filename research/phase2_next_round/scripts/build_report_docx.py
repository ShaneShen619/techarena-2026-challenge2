"""Build the final Chinese Word report from frozen Markdown and CSV artifacts."""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

TASK = Path(__file__).resolve().parents[1]
OUT = TASK / "outputs"
SOURCE = OUT / "LFP_SOH_下一轮逐路线验证与决策报告.md"
TARGET = OUT / "LFP_SOH_下一轮逐路线验证与决策报告.docx"
NAVY = "17324D"
TEAL = "0A7E8C"
ORANGE = "D97B29"
LIGHT = "EAF1F5"
PALE = "F5F8FA"
WHITE = "FFFFFF"
GRAY = "5A6872"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, latin="Arial Unicode MS", east="Arial Unicode MS", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clean_inline(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def set_repeat_table_header(table):
    set_repeat_header(table.rows[0])
    for cell in table.rows[0].cells:
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=8.5, bold=True, color=WHITE)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_i, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=8.2)
        if row_i and row_i % 2 == 0:
            for cell in row.cells:
                shade(cell, PALE)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=8, color=GRAY)


def add_callout(doc, title, text, color=TEAL):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.35)
    table.columns[1].width = Cm(16.6)
    shade(table.cell(0, 0), color)
    shade(table.cell(0, 1), LIGHT)
    table.cell(0, 0).text = ""
    p = table.cell(0, 1).paragraphs[0]
    r = p.add_run(title + "\n")
    set_run_font(r, size=11, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=GRAY)
    set_cell_margin(table.cell(0, 1), 140, 180, 140, 180)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.7)
section.different_first_page_header_footer = True

# Global styles.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial Unicode MS"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 12), ("Subtitle", 13, TEAL, 0, 8),
    ("Heading 1", 18, NAVY, 16, 8), ("Heading 2", 13, TEAL, 12, 5),
    ("Heading 3", 11, ORANGE, 9, 4),
]:
    style = styles[name]
    style.font.name = "Heiti SC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Header/footer.
header = section.header.paragraphs[0]
header.text = "LFP SOH 第二阶段  ·  下一轮逐路线验证与决策"
for run in header.runs:
    set_run_font(run, size=8.5, bold=True, color=NAVY)
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_field(section.footer.paragraphs[0])

# Cover.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(35)
r = p.add_run("LFP SOH  /  PHASE 2")
set_run_font(r, size=11, bold=True, color=TEAL)
p = doc.add_paragraph(style="Title")
p.add_run("下一轮逐路线验证\n与决策报告")
p = doc.add_paragraph(style="Subtitle")
p.add_run("从“模型分数”回到“容量证据”：R0–R6 完整闭环")

doc.add_paragraph().paragraph_format.space_after = Pt(8)
status = doc.add_table(rows=2, cols=2)
status.autofit = False
items = [
    ("研究执行", "R0–R6 已完成"), ("性能高目标", "未达到"),
    ("独立四串确认", "尚无合格标签"), ("官方 CK1–CK7 精度", "未验证"),
]
for cell, (label, value) in zip([c for row in status.rows for c in row.cells], items):
    shade(cell, LIGHT)
    set_cell_margin(cell, 180, 180, 180, 180)
    p = cell.paragraphs[0]
    r = p.add_run(label + "\n")
    set_run_font(r, size=8.5, bold=True, color=TEAL)
    r = p.add_run(value)
    set_run_font(r, size=12, bold=True, color=NAVY)
status.alignment = WD_TABLE_ALIGNMENT.CENTER

doc.add_paragraph().paragraph_format.space_after = Pt(8)
add_callout(doc, "核心判断",
            "冻结候选仍为 2.851 pp；严格去温度在外层面板改善平均但恶化最差芯与最大误差，温度是平均—尾部取舍。电压波形容量增量未成立；CK1–CK7 真值隐藏。")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(44)
r = p.add_run("版本 2026-09-29  ·  Europe/Berlin\n")
set_run_font(r, size=10, bold=True, color=NAVY)
r = p.add_run("项目目录：Arena 阶段 2 官方材料 / research / phase2_next_round")
set_run_font(r, size=8.5, color=GRAY)
p.add_run().add_break(WD_BREAK.PAGE)

# Document control and manual contents.
doc.add_heading("文档控制", level=1)
control = doc.add_table(rows=5, cols=2)
control_data = [
    ("用途", "团队技术决策、复现和下一轮测量交接"),
    ("主证据", "D1 v1.5 六芯 180 目标；D3 CK0–CK7 严格前缀"),
    ("选择规则", "全信号 TCN 为冻结候选；严格去温度与尾部加权仅作探索对照"),
    ("不作出的声明", "不将 D1、接口通过、电压预测或合成结果写成官方容量精度"),
    ("配套文件", "逐点 CSV、验收 JSON、测量计划、候选压缩包与复现命令"),
]
for row, (a, b) in zip(control.rows, control_data):
    row.cells[0].text, row.cells[1].text = a, b
    shade(row.cells[0], NAVY)
    for run in row.cells[0].paragraphs[0].runs:
        set_run_font(run, size=9, bold=True, color=WHITE)
    for run in row.cells[1].paragraphs[0].runs:
        set_run_font(run, size=9)
    set_cell_margin(row.cells[0]); set_cell_margin(row.cells[1])
control.style = "Table Grid"

doc.add_heading("目录", level=1)
for item in ["1 项目问题与本轮目标", "2 证据等级与协议", "3 R0：数据与环境门槛",
             "4 R1：复杂模型到底利用了什么", "5 R2：浅充窗口的反事实检验",
             "6 R3：脉冲、恢复与温度匹配", "7 R4：容量、SOC 与四串截止",
             "8 R5：外部资格和新增测量", "9 R6：候选、门控与官方回放",
             "10 验收结论", "11 剩余不确定性与下一步", "12 复现与来源",
             "附录 A 数据资格", "附录 B 实验索引"]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)
p.add_run().add_break(WD_BREAK.PAGE)

doc.add_heading("执行摘要", level=1)
for paragraph in (OUT / "EXECUTIVE_SUMMARY.md").read_text().split("\n\n")[1:]:
    if paragraph.strip():
        doc.add_paragraph(clean_inline(paragraph.replace("\n", " ")))
add_callout(doc, "决策",
            "安全备选：CK0 hold。探索候选：多温浅充回退，仅研究使用。下一次投入优先用于全新四串同协议容量标签，而不是继续在六芯开发面板上扫模型。",
            color=ORANGE)
doc.add_picture(str(OUT / "figures" / "d1_targets.png"), width=Cm(16.8))
cap = doc.add_paragraph("图 1  D1 v1.5 指标与预注册绝对目标", style=None)
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in cap.runs: set_run_font(run, size=8.5, color=GRAY)

# Parse the detailed Markdown, skipping its title and metadata preface.
lines = SOURCE.read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if line.startswith("# ") or line.startswith("**版本") or line.startswith("**范围") or line.startswith("**结论状态"):
        i += 1
        continue
    if line.startswith("## "):
        heading = clean_inline(line[3:])
        doc.add_heading(heading, level=1)
        if heading.startswith("5. R2"):
            pass
        if heading.startswith("9. R6"):
            doc.add_picture(str(OUT / "figures" / "official_prefix.png"), width=Cm(16.4))
            cap = doc.add_paragraph("图 3  官方前缀回放；灰区真值隐藏", style=None)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs: set_run_font(run, size=8.5, color=GRAY)
        i += 1
        continue
    if line.startswith("### "):
        doc.add_heading(clean_inline(line[4:]), level=2)
        i += 1
        continue
    if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace("-", "").replace(":", "").strip()) == set():
        headers = [clean_inline(x) for x in line.strip("|").split("|")]
        data = []
        i += 2
        while i < len(lines) and lines[i].startswith("|"):
            data.append([clean_inline(x) for x in lines[i].strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(headers))
        for j, value in enumerate(headers): table.cell(0, j).text = value
        for values in data:
            cells = table.add_row().cells
            for j, value in enumerate(values): cells[j].text = value
        set_repeat_table_header(table)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(clean_inline(line[2:]))
        i += 1
        continue
    if line.strip():
        parts = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "- ")):
            parts.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph(clean_inline(" ".join(parts)))
        if "R1" in p.text and "五组实验" in p.text:
            pass
        continue
    i += 1

# Append evidence inventory and experiment index.
doc.add_page_break()
doc.add_heading("附录 A：数据资格清单", level=1)
elig = pd.read_csv(OUT / "data_eligibility.csv")
show_cols = ["source", "physical_entities", "topology", "label_or_target", "evidence_tier", "qualification"]
table = doc.add_table(rows=1, cols=len(show_cols))
for j, col in enumerate(show_cols): table.cell(0, j).text = col
for row in elig[show_cols].itertuples(index=False):
    cells = table.add_row().cells
    for j, value in enumerate(row): cells[j].text = str(value)
set_repeat_table_header(table)

doc.add_page_break()
appendix_b = doc.add_paragraph(style="Heading 1")
appendix_b.paragraph_format.left_indent = Cm(0)
appendix_b.paragraph_format.first_line_indent = Cm(0)
appendix_b.add_run("附录 B：实验索引")
idx = pd.read_csv(TASK / "notes" / "EXPERIMENT_INDEX.csv")
table = doc.add_table(rows=1, cols=4)
for j, col in enumerate(idx.columns): table.cell(0, j).text = col
for row in idx.itertuples(index=False):
    cells = table.add_row().cells
    for j, value in enumerate(row): cells[j].text = str(value)
set_repeat_table_header(table)

doc.add_heading("附录 C：交付物定位", level=1)
for path, purpose in [
    ("outputs/acceptance_results.json", "机器可读验收"),
    ("outputs/method_comparison.csv", "D1 统一方法比较"),
    ("outputs/predictions_long.csv", "R1/R2/D3 逐点结果"),
    ("outputs/official_prefix_diagnostics.csv", "CK0–CK7 严格前缀与分支"),
    ("outputs/measurement_plan.md", "独立四串采集和开封流程"),
    ("outputs/measurement_schema.csv", "采集字段模板"),
    ("outputs/candidate_package_manifest.json", "候选压缩包哈希"),
    ("REPRODUCE.md", "复现入口"),
]:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(path)
    set_run_font(r, bold=True, color=NAVY)
    p.add_run(" — " + purpose)

# Add the ablation figure near the appendices as a visual reference.
doc.add_picture(str(OUT / "figures" / "r1_ablations.png"), width=Cm(16.4))
cap = doc.add_paragraph("图 2  R1 配对消融：温度呈平均—尾部取舍；电压容量增量很弱")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in cap.runs: set_run_font(run, size=8.5, color=GRAY)

# Core properties.
doc.core_properties.title = "LFP SOH 下一轮逐路线验证与决策报告"
doc.core_properties.subject = "R0–R6 strict-prefix research, signal attribution and measurement decision"
doc.core_properties.author = "LFP SOH Phase 2 research team"
doc.core_properties.keywords = "LFP, SOH, capacity, causal prefix, ablation, 4S"

doc.save(TARGET)
print(TARGET)
