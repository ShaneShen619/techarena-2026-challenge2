"""Build the Chinese technical report from the reviewed Markdown source."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TASK = Path(__file__).resolve().parents[1]
SOURCE = TASK / "outputs/方案A_温度修复与提升报告.md"
OUTPUT = TASK / "outputs/方案A_温度修复与提升报告.docx"
FONT = "PingFang SC"


def font(style, size: float, bold: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for key in ("ascii", "hAnsi", "eastAsia"):
        rf.set(qn(f"w:{key}"), FONT)


def set_run_font(run) -> None:
    run.font.name = FONT
    if run._element.rPr is not None and run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def add_inline(paragraph, value: str) -> None:
    # Bold and inline code survive the Markdown-to-Word conversion.
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", value)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1]); run.font.size = Pt(8.5)
        else:
            run = paragraph.add_run(token)
        set_run_font(run)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    box = tcpr.find(qn("w:tcBorders"))
    if box is None:
        box = OxmlElement("w:tcBorders"); tcpr.append(box)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:color"), "D9D9D9")
        el.set(qn("w:sz"), "4")
        box.append(el)


def set_cell_margin(cell, top=110, start=95, bottom=110, end=95) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), str(amount))
        item.set(qn("w:type"), "dxa")
        mar.append(item)
    tcpr.append(mar)


def add_table(doc, block: list[str]) -> None:
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in block]
    rows = [rows[0]] + rows[2:]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    widths = [Cm(8.0), Cm(2.35), Cm(2.4), Cm(2.1), Cm(2.45)]
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell); set_cell_margin(cell)
            if ri == 0:
                shade(cell, "E9EFF5")
            elif ri % 2 == 0:
                shade(cell, "F7F9FB")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value.replace("**", ""))
            run.bold = ri == 0 or "**" in value
            run.font.size = Pt(8.4)
            set_run_font(run)
        if ri == 0:
            trpr = table.rows[ri]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.05); section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.55); section.right_margin = Cm(2.45)
    section.header_distance = Cm(.9); section.footer_distance = Cm(.9)
    normal = doc.styles["Normal"]
    font(normal, 9.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    title = doc.styles["Title"]
    font(title, 19, True)
    title.paragraph_format.space_after = Pt(12)
    if title.element.pPr is not None:
        old_border = title.element.pPr.find(qn("w:pBdr"))
        if old_border is not None:
            title.element.pPr.remove(old_border)
    h1 = doc.styles["Heading 1"]
    font(h1, 13, True)
    h1.paragraph_format.space_before = Pt(15)
    h1.paragraph_format.space_after = Pt(7)
    h1.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    font(caption, 8.5)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    header = section.header.paragraphs[0]
    header.text = "TechArena 2026  LFP 电池 SOH 第二阶段"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8); run.font.color.rgb = RGBColor(90, 90, 90); set_run_font(run)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("项目内部技术报告  ·  ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8); set_run_font(run)

    lines = SOURCE.read_text().splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].strip()
        if not raw:
            i += 1; continue
        if raw.startswith("# "):
            title_text = raw[2:].replace("方案 A ", "方案 A ")
            add_inline(doc.add_paragraph(style="Title"), title_text)
        elif raw.startswith("## "):
            heading = re.sub(r"[^\w\s一-龥]+", "", raw[3:]).strip()
            paragraph = doc.add_paragraph(style="Heading 1")
            if heading.startswith("7 "):
                paragraph.paragraph_format.page_break_before = True
            add_inline(paragraph, heading)
        elif raw.startswith("!["):
            match = re.match(r"!\[(.+)\]\(([^)]+)\)", raw)
            if match:
                path = SOURCE.parent / match.group(2)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                width = Cm(10.5) if "mechanism_truth_vs_fallback" in path.name else Cm(13.6)
                p.add_run().add_picture(str(path), width=width)
                cp = doc.add_paragraph(style="Caption")
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(cp, match.group(1))
        elif raw.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block)
            continue
        elif re.match(r"\d+\.\s", raw):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(.45)
            add_inline(p, raw)
        else:
            add_inline(doc.add_paragraph(), raw)
        i += 1
    doc.core_properties.title = "方案 A 温度修复与容量映射提升技术报告"
    doc.core_properties.subject = "TechArena 2026 Challenge 2 LFP 电池 SOH"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
