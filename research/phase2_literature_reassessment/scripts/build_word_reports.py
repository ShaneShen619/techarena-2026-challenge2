"""Deterministic Chinese Word exports from the approved Markdown source files."""
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TASK=Path(__file__).resolve().parents[1];OUT=TASK/'outputs'
FILES=[
 ('LFP_SOH_第二阶段文献综述与技术路线_深度重评','TechArena 2026 第二阶段 LFP 容量 SOH 文献证据重评与技术路线','文献综述与技术决策报告'),
 ('LFP_SOH_核心论文精读与证据附册','LFP SOH 第二阶段核心论文精读与证据附册','原始证据卡汇编'),
]
NAVY=RGBColor(25,48,75);GRAY=RGBColor(75,84,96)
def font(style,name='PingFang SC',size=10.5,bold=False,color=None):
 style.font.name=name;style.font.size=Pt(size);style.font.bold=bold
 style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)
 if color:style.font.color.rgb=color
def clean(s):
 s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1（\2）',s)
 return s.replace('**','').replace('`','').strip()
def add_page(paragraph):
 paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 paragraph.add_run('第 ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');paragraph._p.append(fld)
 paragraph.add_run(' 页')
def build(stem,title,sub):
 source=(OUT/(stem+'.md')).read_text(encoding='utf-8')
 doc=Document();sec=doc.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
 sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
 sec.header_distance=Inches(.35);sec.footer_distance=Inches(.35)
 styles=doc.styles
 font(styles['Normal'],size=10.5)
 styles['Normal'].paragraph_format.space_after=Pt(5)
 styles['Normal'].paragraph_format.line_spacing=1.28
 for n,sz,bef,aft in [('Title',20,0,12),('Heading 1',14,15,7),('Heading 2',12,11,5),('Heading 3',10.8,9,4),('Heading 4',10.5,7,3)]:
  font(styles[n],size=sz,bold=True,color=NAVY)
  styles[n].paragraph_format.space_before=Pt(bef);styles[n].paragraph_format.space_after=Pt(aft)
  styles[n].paragraph_format.keep_with_next=True
 title_ppr=styles['Title']._element.get_or_add_pPr()
 for child in list(title_ppr):
  if child.tag==qn('w:pBdr'):title_ppr.remove(child)
 font(styles['List Bullet'],size=10.2)
 styles['List Bullet'].paragraph_format.space_after=Pt(3)
 header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 rr=header.add_run('LFP SOH  第二阶段文献重评  ·  2026-09-29');rr.font.size=Pt(8);rr.font.color.rgb=GRAY
 add_page(sec.footer.paragraphs[0])
 p=doc.add_paragraph(style='Title');p.add_run(title)
 ppr=p._p.get_or_add_pPr()
 for child in list(ppr):
  if child.tag==qn('w:pBdr'):ppr.remove(child)
 p=doc.add_paragraph();p.add_run(sub+'  |  2026-09-29  |  研究证据与实施参考').italic=True
 p.paragraph_format.space_after=Pt(12)
 in_code=False
 for raw in source.splitlines()[1:]:
  s=raw.rstrip()
  if s.startswith('```'):
   in_code=not in_code;continue
  if not s.strip():continue
  if in_code:
   p=doc.add_paragraph(style='Normal');p.paragraph_format.left_indent=Inches(.18)
   r=p.add_run(s);r.font.name='Menlo';r.font.size=Pt(8.2)
   continue
  m=re.match(r'^(#{1,6})\s+(.*)$',s)
  if m:
   level=min(len(m.group(1)),4)
   if stem.startswith('LFP_SOH_核心'):
    # Source card sections are nested three levels under each card.
    if level==2:level=1
    elif level>2:level=min(level-1,4)
   doc.add_heading(clean(m.group(2)),level=level);continue
  if s.startswith('- '):
   doc.add_paragraph(clean(s[2:]),style='List Bullet');continue
  p=doc.add_paragraph(style='Normal');p.add_run(clean(s))
 target=OUT/(stem+'.docx');doc.save(target)
 print(target.name,target.stat().st_size)
for x in FILES:build(*x)
