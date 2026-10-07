from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
md=(root/'outputs/LFP_SOH_第二阶段文献综述与技术路线.md').read_text(encoding='utf-8')
output=root/'outputs/LFP_SOH_第二阶段文献综述与技术路线.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
sec.header_distance=Cm(0.8);sec.footer_distance=Cm(0.8)

def font(st,size,bold=False):
 st.font.name='Songti SC';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
 rPr=st.element.get_or_add_rPr();rFonts=rPr.rFonts
 if rFonts is None: rFonts=OxmlElement('w:rFonts');rPr.insert(0,rFonts)
 rFonts.set(qn('w:eastAsia'),'Songti SC')
 rFonts.set(qn('w:ascii'),'Songti SC')
 rFonts.set(qn('w:hAnsi'),'Songti SC')

styles=doc.styles
font(styles['Normal'],10.2)
styles['Normal'].paragraph_format.space_after=Pt(5.5)
styles['Normal'].paragraph_format.line_spacing=Pt(16.2)
for name,size,before,after in [('Title',18,0,12),('Heading 1',14,16,7),('Heading 2',12,12,5),('Heading 3',10.8,8,4)]:
 st=styles[name];font(st,size,True)
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
 st.paragraph_format.keep_with_next=True
 st.paragraph_format.line_spacing=Pt(size+5)
 # No inherited title/heading bottom border or theme color
 ppr=st.element.get_or_add_pPr();pbdr=ppr.find(qn('w:pBdr'))
 if pbdr is not None:ppr.remove(pbdr)

for name,size in [('Small Table',8.7),('Code',8.5),('Reference',9.4)]:
 if name not in [s.name for s in styles]: styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st=styles[name];font(st,size);st.paragraph_format.space_after=Pt(2 if name=='Reference' else 3);st.paragraph_format.line_spacing=Pt(size+3.5 if name=='Reference' else size+4)

# Page number, restrained footer.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.style='Small Table'
run=footer.add_run('LFP SOH 第二阶段文献综述与技术路线   ·   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def add_inline(p,s):
 s=s.replace('`','')
 parts=re.split(r'(\*\*.*?\*\*)',s)
 for part in parts:
  if not part:continue
  r=p.add_run(part[2:-2] if part.startswith('**') and part.endswith('**') else part)
  if part.startswith('**') and part.endswith('**'):r.bold=True
  r.font.name='Songti SC'
  rPr=r._element.get_or_add_rPr();rf=rPr.rFonts
  if rf is None:rf=OxmlElement('w:rFonts');rPr.insert(0,rf)
  rf.set(qn('w:eastAsia'),'Songti SC')

lines=md.splitlines();i=0
while i<len(lines):
 line=lines[i].rstrip()
 if not line.strip():i+=1;continue
 if line.startswith('```'):
  i+=1
  while i<len(lines) and not lines[i].startswith('```'):
   p=doc.add_paragraph(style='Code');p.paragraph_format.left_indent=Cm(.45);add_inline(p,lines[i]);i+=1
  i+=1;continue
 if line.startswith('|') and i+1<len(lines) and re.match(r'^\|\s*[-:| ]+\|$',lines[i+1]):
  headers=[v.strip() for v in line.strip('|').split('|')]
  i+=2;rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   rows.append([v.strip() for v in lines[i].strip('|').split('|')]);i+=1
  table=doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  if len(headers)==4:
   widths=[Cm(x) for x in ([3.0,4.5,5.1,4.4] if '方法族' in headers[0] else [2.4,5.0,5.1,4.5])]
  else:widths=[Cm(17/len(headers)) for _ in headers]
  for j,h in enumerate(headers):
   c=table.rows[0].cells[j];c.width=widths[j];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.style='Small Table';rr=p.add_run(h);rr.bold=True
   tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF0');tcPr.append(sh)
  hdr=OxmlElement('w:tblHeader');hdr.set(qn('w:val'),'true');table.rows[0]._tr.get_or_add_trPr().append(hdr)
  for k,row in enumerate(rows):
   cells=table.add_row().cells
   for j,val in enumerate(row):
    cells[j].width=widths[j];cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    p=cells[j].paragraphs[0];p.style='Small Table';add_inline(p,val)
    if k%2:tcPr=cells[j]._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F7F8F9');tcPr.append(sh)
   cant=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(cant)
  doc.add_paragraph().paragraph_format.space_after=Pt(0)
  continue
 if line.startswith('# '):
  p=doc.add_paragraph(style='Title');add_inline(p,line[2:]);i+=1;continue
 if line.startswith('## '):
  p=doc.add_paragraph(style='Heading 1');add_inline(p,line[3:]);i+=1;continue
 if line.startswith('### '):
  p=doc.add_paragraph(style='Heading 2');add_inline(p,line[4:]);i+=1;continue
 if line.startswith('- '):
  p=doc.add_paragraph(style='List Bullet');add_inline(p,line[2:]);i+=1;continue
 is_ref=bool(re.match(r'^\[[A-Za-z][A-Za-z0-9]+\]',line))
 p=doc.add_paragraph(style='Reference' if is_ref else 'Normal')
 add_inline(p,line.replace('*','') if is_ref else line)
 i+=1

# Core document metadata
props=doc.core_properties;props.title='LFP SOH 第二阶段文献综述与技术路线';props.subject='TechArena 2026 Topic 1 Challenge 2 技术决策';props.author='项目研究组'
doc.save(output)
print(output)
