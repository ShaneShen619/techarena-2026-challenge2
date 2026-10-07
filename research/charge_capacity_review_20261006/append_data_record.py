"""Append verified operating-capacity findings to the existing desktop DOCX."""
from pathlib import Path
from copy import deepcopy
import csv
import hashlib
import json
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
TARGET = Path('/Users/shane/Desktop/官方原始数据分析记录.docx')
BACKUP = BASE / '官方原始数据分析记录_更新前_v1.0.docx'
if not BACKUP.exists():
    shutil.copy2(TARGET, BACKUP)
original_digest = hashlib.sha256(TARGET.read_bytes()).hexdigest()
doc = Document(BACKUP)
result = json.loads((BASE / 'endpoint_summary.json').read_text())
findings = json.loads((BASE / 'findings.json').read_text())
with (BASE / 'fixed_CC_voltage_windows.csv').open() as f:
    windows = list(csv.DictReader(f))
deep = [r for r in windows if float(r['V_lo']) == 12.2 and int(r['setpoint_C']) == 25]
assert len(deep) == 7
guide = next(p for p in doc.paragraphs if p.text == '后续章节的记录方式')
original_body = list(doc._element.body)

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    for n, part in enumerate(text.split('**')):
        r = p.add_run(part)
        r.bold = bool(n % 2)
    return p

def heading(text, level=2, new_page=False):
    p = doc.add_heading(text, level=level)
    if new_page:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), value)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for j, value in enumerate(headers):
        t.rows[0].cells[j].text = str(value)
    for values in rows:
        cells = t.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = str(value)
    for n, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if n == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '85'), ('bottom', '85'), ('left', '110'), ('right', '110')]:
                e = OxmlElement('w:' + side)
                e.set(qn('w:w'), value)
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            tcpr.append(margins)
            sh = OxmlElement('w:shd')
            sh.set(qn('w:fill'), '173B50' if n == 0 else ('F1F5F7' if n % 2 == 0 else 'FFFFFF'))
            tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.bold = n == 0
                    r.font.color.rgb = RGBColor(255, 255, 255) if n == 0 else RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def hyperlink(label, target):
    p = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); pr.append(color)
    r.append(pr)
    text = OxmlElement('w:t'); text.text = label; r.append(text)
    link.append(r); p._p.append(link)
    return p

heading('第二章 充放电量与容量相关趋势', 1, True)
para('本章继续检查第二阶段官方四串电池组的充放电量、事件起止电压及其随时间的变化。范围为 data/operation/ 全部 13 个已公开文件，共 **1,572,894 行**，2025 年 1 月 19 日至 9 月 11 日，另核对 CK0 参考放电。')
heading('主要结论')
para('充电结束电压非常稳定，但充电起点分为常规补充充电和较深充电两类。**7 次 25°C 档位的较深充电，在组压 12.2→13.9 V 的统一窗口内，充入电量由 94.99 Ah 逐次降至 89.95 Ah，首末下降 5.31%。**相邻大电压窗口的下降约为 5.27% 至 5.39%，支持容量相关充电特征存在衰减趋势。')
para('该下降是实测窗口电量的变化，尚不能精确等同于官方参考容量或 SOH 下降。常规充电也出现小幅下降；半小时放电脉冲的电量则由固定电流和时间决定，基本不变。')
heading('数据来源与容量定义')
para('比赛第一阶段关注单芯长期循环老化与寿命，其原始文件位于 dataset original/；第二阶段关注四串电池组在指定日期的容量健康状态。本章主体只统计第二阶段，两者不是同一只电池的前后两个使用时期，测试协议也不能直接互换。')
table(['字段或测量', '含义及容量判断'], [
    ['charge_Ah_cum / discharge_Ah_cum', '试验运行开始后累计充入／放出 Ah，不清零。主要反映电荷通量，不能直接当作总容量。'],
    ['CK0 的 discharged_Ah', '单次参考放电从开始累计的 Ah。曲线末点为 100.412 Ah，官方发布容量为 100.41 Ah。'],
    ['官方容量与 SOH', '满充后以 5.1 A 放电至组端首次 11.2 V 的放出量；SOH = 容量 ÷ 102 Ah × 100%。CK0 为 98.44%。']
], [6.0, 10.8])
para('四枚 102 Ah 电芯串联增加电压，容量分母仍为 102 Ah。运行记录的最低组压为 **11.473 V**，达到或低于 11.2 V 的行数为 **0**，因此这些可见运行段没有提供官方参考容量完整放电的截止证据。CK1—CK7 参考容量仍未公开。')
heading('累计电量的通道语义需要复核')
para('最终运行记录的充电累计为 **16,585.075 Ah**，放电累计为 **15,214.406 Ah**，均是长期累计通量。第 3 段虽未提供，官方说明计数器继续累加；月度 checkup 的通量不计入运行计数器。两列相减不能直接给出总容量或当前剩余电量。')
para('2025-01-20 01:06:09 至 01:36:08 的 181 行脉冲中，电流约 −20.4 A，积分放电 **10.19398 Ah**，但 charge_Ah_cum 增加 **10.200 Ah**，discharge_Ah_cum 增加 **0 Ah**。这是局部通道语义与列名不一致的直接证据；不能据此认定两列全程交换。下文电量均从电流与实测时间积分，不以计数器列名直接取差。')

heading('事件识别与起止电压', 2, True)
para('先按同一文件内 |I|>0.5 A 的同向连续段识别事件，要求持续至少 15 分钟，不跨文件，不跨超过 30 秒的间隔，也不跨时间倒序。电流回到近零或改变方向时结束当前事件；重复时间戳的零时长不贡献积分。')
para('充电候选要求末个带载组压 ≥13.95 V 且中位电流 >15 A，共 41 次。其充电前组压 ≥13.3 V 的 33 次归为常规补充充电，低于 13.3 V 的 8 次归为较深充电。两组实际起点范围明显分离；该阈值只用于本次描述分组。41 次充电候选的最大采样间隔均为 10 秒。')
table(['事件类型', '次数', '首个带载\n组压 V', '末个带载\n组压 V', '整段电量\nAh'], [
    ['常规补充充电', '33', '13.407–13.439', '14.008–14.010', '20.637–21.758'],
    ['较深充电', '8', '11.617–12.037', '14.009–14.010', '91.975–100.143'],
    ['约 20.4 A 半小时放电', '41', '13.365–13.602', '13.160–13.226', '10.194–10.200'],
    ['连续动态放电 ≥15 Ah', '444', '13.065–13.400', '12.846–13.305', '15.618–50.232']
], [4.6, 1.2, 3.7, 3.7, 3.6])
para('充电起点必须区分通电前与通电后。33 次常规充电前的近零电流组压为 **13.347–13.363 V**，范围仅 16 mV；通电后的首个组压为 13.407–13.439 V。8 次较深充电前为 **11.570–11.953 V**，范围 383 mV。通电、停流时的电压跳变不能直接解释为 SOC 变化。')
para('41 次充电的末个带载组压为 14.008–14.010 V，范围仅 2 mV；停流后的首个组压回落至 **13.979–13.987 V**。比较事件端点时应始终采用相同的带载／近零电流口径。')
heading('常规补充充电的电量变化')
para('从 33 次常规充电中排除各运行段首个事件，以减少此前工况差异，保留 28 次；按温箱档位分别比较最早 3 次与最晚 3 次的电量中位数，不混合温度。')
table(['温箱档位', '匹配次数', '最早 3 次\n中位 Ah', '最晚 3 次\n中位 Ah', '相对变化'], [
    ['25°C', '9', '21.445', '20.945', '−2.33%'],
    ['45°C', '19', '21.407', '20.703', '−3.29%']
], [3.0, 2.4, 3.8, 3.8, 3.8])
para('25°C 档最早组为 2 月 16 日、3 月 19 日、4 月 22 日，最晚组为 8 月 2 日、8 月 7 日、9 月 6 日；整段平均芯温分别为 26.333°C、26.171°C。45°C 档最早组为 1 月 24 日、1 月 29 日、2 月 3 日，最晚组为 8 月 17 日、8 月 22 日、8 月 27 日；平均芯温分别为 46.458°C、46.203°C。')
para('常规补充充电电量的下降可作为辅助趋势证据，但实际芯温、充电前电压和前序运行状态仍有变化，不能直接将 −2.33% 或 −3.29% 写成官方总容量损失。')

heading('较深充电的固定电压窗口趋势', 2, True)
para('8 次较深充电中，7 次在 25°C 温箱档位发生，唯一一次 45°C 事件不参与这条趋势。由于整段充电的低压起点有浮动，本次将各事件统一截取为组压 **12.2→13.9 V**，以提高事件可比性。')
para('取每次充电首次向上跨越两个阈值的时刻，并在相邻样本间线性插值。窗口内每个电流样本及插值边界均要求在 20.2–20.6 A 之间，实际约 20.4 A；再按时间和电流做梯形积分。整段事件平均芯温为 **26.33–26.59°C**，不是人为将实测温度统一为 25°C。')
table(['充电开始日期', '运行段', '固定窗口\n电量 Ah', '整段平均\n芯温 °C'], [
    [r['t_start'][:10], r['segment'], f"{float(r['Q_window_Ah']):.2f}", f"{float(r['mean_temp_C']):.2f}"] for r in deep
], [4.6, 2.2, 5.1, 4.9])
para('7 次窗口电量逐次下降。首末事件比较为 **94.99494→89.95058 Ah**，减少 **5.04437 Ah，下降 5.31%**，对应 2025-02-11 至 2025-09-01 的跨度。前 3 次与后 3 次中位数为 **93.61225→90.61169 Ah，下降 3.21%**；这比较的是不同时间范围，不能与首末变化混用。')
heading('相邻大电压窗口的敏感性')
table(['组压窗口 V', '最早电量 Ah', '最晚电量 Ah', '首末变化'], [
    [f"{r['V_lo']:.1f}→{r['V_hi']:.1f}", f"{r['first_Ah']:.2f}", f"{r['last_Ah']:.2f}", f"{r['first_to_last_change_pct']:.2f}%"] for r in result['deep_window_comparisons']
], [4.6, 4.1, 4.1, 4.0])
para('四个窗口的 7 次电量均逐次减少，首末变化为 **−5.27% 至 −5.39%**，说明本组较深充电的下降对所检查的大窗口较稳定。它们共用同一批事件，不能当作四次独立验证。')
para('本章据此记录：**在相近温度、相同电流和固定组压窗口下，容量相关的充电电量特征持续下降。**若要换算为总容量或 SOH，仍需确认电压对应的 SOC 范围和老化后的曲线变化，并用同组参考放电容量标定。')

heading('放电事件和窄窗口的使用边界', 2, True)
para('41 次半小时脉冲电流约 −20.4 A，持续 29.983–30.000 分钟，放出 10.194–10.200 Ah。这个电量由程序设定的电流和持续时间决定，基本不随容量变化。起止组压存在浮动，更适合研究同温度、相似初始状态下的电压响应。')
table(['脉冲温箱档位', '次数', '首个带载组压 V', '末个带载组压 V'], [
    ['25°C', '16', '13.365–13.411', '13.160–13.171'],
    ['45°C', '25', '13.462–13.602', '13.212–13.226']
], [4.6, 2.2, 5.0, 5.0])
para('另有 444 个连续动态放电段满足整段电量 ≥15 Ah；其带载端点、电流和持续时间范围较大。它们是按电流方向分出的操作片段，不等同于 444 次完整循环，也不能直接比较整段电量得出容量衰减。')
heading('窄电压窗口的电量下降并不等于容量同比下降')
para('同一批 19 次匹配的 45°C 常规充电，按最早 3 次与最晚 3 次中位数比较：13.5→13.9 V 窗口变化为 **−4.51%**，13.6→13.9 V 为 **−12.44%**，13.7→13.9 V 为 **−4.67%**，13.8→13.9 V 约 **−0.02%**。局部电量对窗口选择明显敏感，不能挑选某个百分比当作总容量下降。')
para('LFP 的电压平台、滞后、内阻和极化，以及实际温度和前序充放电状态，会改变相同端电压所对应的 SOC。较深充电的大窗口结果具有更一致的趋势，但仍只代表本组数据中的容量相关代理指标。')
heading('积分口径与因果边界')
para('连续充电电量按 Σ[(I前+I后)×Δt]/7200 计算，电流单位 A、时间单位秒；放电取相反数。整段电量只积分首末带载样本，不向步骤边界外推，不跨长缺口。固定窗口在阈值处插值，只积分窗口内电流。')
para('大窗口电量与“20.4 A×实测窗口时长”的最大相对差约 **0.00136%**，单位换算与恒流时长核对一致。全部分析覆盖 CK7 之后的数据；用于早期检查点预测时必须按当时可见历史重算。9 月事件不能用于回填 8 月或更早检查点的特征。')
para('本次没有新增官方容量标签，也没有验证容量预测误差。后续优先保留较深充电大窗口电量、常规补充充电电量、实际温度和端点状态，并用配对参考容量检查其与 SOH 的关系。')

heading('本章证据与复现入口', 2, True)
para('以下文件位于项目文件夹“Arena 阶段 2 官方材料”。逐事件记录保留开始前、首末带载、停流后电压，以及温度、电流、电量、时间和采样间隔，便于复核筛选条件。')
for label, path in [
    ('充放电字段与 CK0 核对结果', BASE / 'findings.json'),
    ('充放电端点与分组汇总', BASE / 'endpoint_summary.json'),
    ('逐事件电压电流温度与电量', BASE / 'operating_events.csv'),
    ('固定恒流电压窗口的逐次电量', BASE / 'fixed_CC_voltage_windows.csv'),
    ('事件端点与窗口积分复算脚本', BASE / 'analyze_event_endpoints.py'),
    ('既有 41 次放电脉冲计数器核对', ROOT / 'research/phase2_baseline_repair/outputs/counter_audit.csv'),
    ('官方数据字段与容量定义', ROOT / 'data/DATA_DESCRIPTION.md'),
]:
    hyperlink(label, path.resolve().as_uri())
hyperlink('LFP 电压平台与温度相关滞后研究', 'https://publications.rwth-aachen.de/record/1019597/files/1019597.pdf')
para('本章补充说明第一阶段字段差异：抽查单芯原始文件的 step_capacity_Ah 是当前步骤累计量，第 2 循环约 100 A 放至 2.5 V 的末值为 105.1542 Ah，独立积分约 105.1702 Ah。确认此前满充和步骤完整后可反映该工况下的单芯放电容量，但不能直接替代第二阶段 5.1 A、四串组压 11.2 V 的参考容量。本章没有审核第一阶段全部文件或循环。')
para('章节新增于 2026 年 10 月 6 日。41 次完整充电候选与既有事件审查数量一致；本次从原始运行文件重新提取整段事件，采用的边界与官方示例 13.4 V 起始局部窗口不同，因此两个电量数值不应直接混用。')

# Place the new chapter before the existing recording guide without rebuilding old content.
new_nodes = [n for n in list(doc._element.body) if n not in original_body]
for node in new_nodes:
    guide._p.addprevious(node)
doc.paragraphs[1].text = 'TechArena 2026 Challenge 2　｜　版本 1.1　｜　更新于 2026 年 10 月 6 日'
doc.paragraphs[2].text = '本记录用于持续汇总官方原始数据的字段含义、分布规律、缺失与异常，以及对后续分析的影响。当前已完成温度章节和充放电量与容量相关趋势章节；后续字段总结可作为独立章节追加，并在更新记录中登记范围和结论变化。'
update = doc.tables[0]
new_row = deepcopy(update.rows[-1]._tr)
update._tbl.append(new_row)
for cell, text in zip(update.rows[-1].cells, ['2026-10-06', '1.1', '新增充放电累计量语义、事件端点、固定电压窗口趋势与使用边界']):
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.size = Pt(9.5)
guide.paragraph_format.keep_with_next = True
next_p = next(p for p in doc.paragraphs if p.text.startswith('后续原始数据总结沿用'))
next_p.text = '后续原始数据总结沿用 Word 的“标题 1”添加第三章、第四章等，章节内部使用“标题 2”组织内容，以便通过导航窗格查阅。当前不预填未分析字段的结论。'
doc.core_properties.keywords = 'LFP SOH 温度 充放电量 固定电压窗口 原始数据 数据质量'
candidate = BASE / '官方原始数据分析记录_候选_v1.1.docx'
doc.save(candidate)
assert hashlib.sha256(TARGET.read_bytes()).hexdigest() == original_digest, 'Desktop document changed during editing; review before replacement.'
shutil.copy2(candidate, TARGET)
print('Updated:', TARGET)
print('Preserved backup:', BACKUP)
print('Paragraphs:', len(doc.paragraphs), 'Tables:', len(doc.tables))
