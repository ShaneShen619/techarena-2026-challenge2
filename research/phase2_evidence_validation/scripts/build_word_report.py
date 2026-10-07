"""Build the Chinese deliverable from the reviewed Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TASK=Path(__file__).resolve().parents[1];OUT=TASK/'outputs'
src=OUT/'LFP_SOH_文献驱动验证与双路线决策报告.md'
dst=OUT/'LFP_SOH_文献驱动验证与双路线决策报告.docx'
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
styles=d.styles
normal=styles['Normal'];normal.font.name='PingFang SC';normal.font.size=Pt(10.2)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC')
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.25
for name,size,color in [('Title',20,RGBColor(25,48,78)),('Heading 1',14,RGBColor(25,65,104)),('Heading 2',11,RGBColor(38,80,110))]:
    s=styles[name];s.font.name='PingFang SC';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=color
    s._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC')
    s.paragraph_format.space_before=Pt(13);s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
# Remove Word's default Title bottom border to avoid the blue line seen in V0 smoke.
for s in (styles['Title'],):
    ppr=s._element.pPr
    if ppr is not None:
        b=ppr.find(qn('w:pBdr'))
        if b is not None:ppr.remove(b)
header=sec.header.paragraphs[0];header.text='LFP SOH 第二阶段 · 文献驱动验证';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(110,120,130)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('2026-09-29  ·  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)
def inline(p,line):
    pieces=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',line)
    for x in pieces:
        if not x:continue
        if x.startswith('**') and x.endswith('**'):
            run=p.add_run(x[2:-2]);run.bold=True
        elif x.startswith('`') and x.endswith('`'):
            run=p.add_run(x[1:-1]);run.font.name='Menlo';run.font.size=Pt(8.5)
        else:p.add_run(x)
def table(lines):
    parsed=[]
    for line in lines:
        parts=[x.strip() for x in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?',x or '') for x in parts):continue
        parsed.append(parts)
    if not parsed:return
    tab=d.add_table(rows=len(parsed),cols=max(map(len,parsed)))
    tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.style='Light Shading Accent 1'
    header_trPr=tab.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');header_trPr.append(repeat)
    for i,row in enumerate(parsed):
        for j,value in enumerate(row):
            cell=tab.cell(i,j);cell.text=re.sub(r'\*\*(.*?)\*\*',r'\1',value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(8.3);run.font.name='PingFang SC';run._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC')
        if i==0:
            for c in tab.rows[0].cells:
                for p in c.paragraphs:
                    for r in p.runs:r.bold=True
    d.add_paragraph()
lines=src.read_text().splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i]);i+=1
        table(block);continue
    m=re.match(r'^!\[(.*?)\]\((.*?)\)$',line)
    if m:
        path=OUT/m.group(2)
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path),width=Cm(15.5))
        cap=d.add_paragraph(m.group(1));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(90,100,110)
        i+=1;continue
    if line.startswith('# '):
        p=d.add_paragraph(style='Title');inline(p,line[2:]);i+=1;continue
    if line.startswith('## '):
        p=d.add_paragraph(style='Heading 1');inline(p,line[3:]);i+=1;continue
    if line.startswith('### '):
        p=d.add_paragraph(style='Heading 2');inline(p,line[4:]);i+=1;continue
    p=d.add_paragraph();inline(p,line);i+=1
d.save(dst)
print(dst)
