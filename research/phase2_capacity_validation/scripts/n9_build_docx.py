"""Build Chinese Word report from audited Markdown source."""
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
T=Path(__file__).resolve().parents[1]
src=T/'outputs/LFP_SOH_真实容量验证与测量计划.md';dest=src.with_suffix('.docx')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.80);sec.right_margin=Inches(.80)
def font(style,size,bold=False):
 style.font.name='PingFang SC';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 rpr=style.element.get_or_add_rPr();f=rpr.rFonts
 if f is None:f=OxmlElement('w:rFonts');rpr.insert(0,f)
 f.set(qn('w:eastAsia'),'PingFang SC')
font(doc.styles['Normal'],9.5);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.16
font(doc.styles['Title'],17,True);doc.styles['Title'].paragraph_format.space_after=Pt(12)
title_ppr=doc.styles['Title'].element.get_or_add_pPr()
for border in title_ppr.findall(qn('w:pBdr')):title_ppr.remove(border)
font(doc.styles['Heading 1'],11.8,True);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(5);doc.styles['Heading 1'].paragraph_format.keep_with_next=True
font(doc.styles['List Bullet'],9.5);doc.styles['List Bullet'].paragraph_format.space_after=Pt(3)
def inline(p,line):
 for c in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))',line):
  if not c:continue
  if c.startswith('**') and c.endswith('**'):p.add_run(c[2:-2]).bold=True
  elif c.startswith('`') and c.endswith('`'):
   r=p.add_run(c[1:-1]);r.font.name='Menlo';r.font.size=Pt(8.2)
  elif c.startswith('[') and '](' in c:
   label,url=c[1:].split('](',1);url=url[:-1]
   r=p.add_run(label+'（'+url+'）');r.font.size=Pt(8.6)
  else:p.add_run(c)
for line in src.read_text().splitlines():
 s=line.strip()
 if not s:continue
 if s.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,s[2:])
 elif s.startswith('## '):p=doc.add_paragraph(style='Heading 1');inline(p,s[3:])
 elif s.startswith('- '):p=doc.add_paragraph(style='List Bullet');inline(p,s[2:])
 else:p=doc.add_paragraph(style='Normal');inline(p,s)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('LFP SOH 阶段二真实容量验证  |  2026 年 9 月 30 日  |  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='LFP SOH 真实容量验证与测量计划'
doc.core_properties.subject='LFP SOH 阶段二数据资格、机制反证和真实容量测量计划'
doc.save(dest);print(dest)
